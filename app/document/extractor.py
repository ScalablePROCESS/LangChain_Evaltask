"""
Extraction de texte depuis des fichiers PDF et DOCX.

Supporte :
- PDF : PyMuPDF (fitz) ou pypdf en fallback
- DOCX : python-docx
- DOC (legacy) : non supporté (prévenir l'utilisateur)

Principes :
- LangChain Backend Engineer : extraction robuste, gestion d'erreurs
- AI Security & Compliance : pas de stockage du fichier, extraction en mémoire
"""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)


def extract_pdf(file_bytes: bytes) -> str:
    """
    Extrait le texte d'un fichier PDF.

    Utilise PyMuPDF (fitz) en priorité (rapide, précis),
    fallback sur pypdf si PyMuPDF n'est pas disponible.
    """
    # Tentative avec PyMuPDF
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                text_parts.append(text.strip())
        doc.close()

        if text_parts:
            return "\n\n".join(text_parts)
        logger.warning("[pdf] Aucun texte extrait (PDF scanné ou vide ?)")
        return ""
    except ImportError:
        logger.debug("[pdf] PyMuPDF non disponible, fallback pypdf")
    except Exception as exc:
        logger.warning("[pdf] PyMuPDF échec: %s, fallback pypdf", exc)

    # Fallback avec pypdf
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text and text.strip():
                text_parts.append(text.strip())

        return "\n\n".join(text_parts)
    except ImportError:
        raise ImportError("Aucune librairie PDF disponible. Installez : pip install pymupdf ou pypdf")
    except Exception as exc:
        logger.error("[pdf] Extraction échouée: %s", exc)
        raise ValueError(f"Extraction PDF échouée : {exc}")


def extract_docx(file_bytes: bytes) -> str:
    """
    Extrait le texte d'un fichier DOCX.
    """
    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(file_bytes))
        text_parts = []

        # Paragraphes
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_parts.append(text)

        # Tableaux
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    text_parts.append(row_text)

        return "\n\n".join(text_parts)
    except ImportError:
        raise ImportError("python-docx n'est pas installé. Installez-le avec : pip install python-docx")
    except Exception as exc:
        logger.error("[docx] Extraction échouée: %s", exc)
        raise ValueError(f"Extraction DOCX échouée : {exc}")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extrait le texte d'un fichier selon son extension.

    Supported:
    - .txt, .md, .csv, .json : UTF-8 direct
    - .pdf : PyMuPDF ou pypdf
    - .docx : python-docx
    """
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in {".txt", ".md", ".csv", ".json"}:
        return file_bytes.decode("utf-8", errors="replace")

    if ext == ".pdf":
        return extract_pdf(file_bytes)

    if ext == ".docx":
        return extract_docx(file_bytes)

    if ext == ".doc":
        raise ValueError("Format .doc (legacy) non supporté. Convertissez en .docx ou .pdf.")

    raise ValueError(f"Format non supporté : {ext}")


# Extensions supportées
SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}
